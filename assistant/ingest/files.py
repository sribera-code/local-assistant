"""Indexation des fichiers locaux : parcours des dossiers et extraction du texte."""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import sqlite3
from pathlib import Path
from typing import Iterator

from .. import db
from ..config import Config, FilesConfig

# Plafond de texte extrait par fichier. Assez haut pour un manuel de 1 400 pages :
# pour des supports de cours, un chapitre coupe est un chapitre introuvable.
MAX_TEXT_CHARS = 5_000_000

# A incrementer a chaque changement de la logique d'extraction : l'empreinte des
# fichiers change, ce qui force leur reextraction a la synchro suivante. Sans cela,
# un fichier deja indexe (meme taille, meme date) garderait l'ancien texte.
EXTRACTOR_VERSION = 2

# pypdf signale chaque police LaTeX qu'il decode imparfaitement : plus d'un millier
# de lignes sur un dossier de cours, qui noieraient la console du serveur.
logging.getLogger("pypdf").setLevel(logging.ERROR)


# ------------------------------------------------------------------ extraction


def _read_text(path: Path) -> str:
    raw = path.read_bytes()[: MAX_TEXT_CHARS * 4]
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # une page illisible ne doit pas perdre tout le document
            continue
        if sum(len(p) for p in pages) > MAX_TEXT_CHARS:
            pages.append(f"\n[... document tronque apres {i + 1} pages ...]")
            break
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    try:
        for sheet in wb.worksheets:
            out.append(f"# Feuille : {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append(" | ".join(cells))
                if len(out) > 5000:
                    out.append("[... feuille tronquee ...]")
                    break
    finally:
        wb.close()
    return "\n".join(out)


def _read_csv(path: Path) -> str:
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        return text
    rows = list(csv.reader(io.StringIO(text), dialect))
    return "\n".join(" | ".join(r) for r in rows[:5000])


_EXTRACTORS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
    ".csv": _read_csv,
}


def extract_text(path: Path) -> str:
    extractor = _EXTRACTORS.get(path.suffix.lower(), _read_text)
    return extractor(path)[:MAX_TEXT_CHARS]


# --------------------------------------------------------------------- parcours


def iter_files(files_cfg: FilesConfig) -> Iterator[Path]:
    extensions = {e.lower() for e in files_cfg.extensions}
    excluded = {d.lower() for d in files_cfg.exclude_dirs}
    max_bytes = files_cfg.max_file_mb * 1024 * 1024

    for root in files_cfg.roots:
        root_path = Path(root).expanduser()
        if not root_path.is_dir():
            continue
        for path in root_path.rglob("*"):
            try:
                if any(part.lower() in excluded for part in path.parts):
                    continue
                if not path.is_file() or path.suffix.lower() not in extensions:
                    continue
                if path.stat().st_size > max_bytes:
                    continue
            except OSError:  # fichier verrouille, lien casse, permission refusee
                continue
            yield path


def stat_hash(path: Path) -> str:
    """Empreinte basee sur taille + date de modification.

    Bien plus rapide que de relire le contenu : suffit pour decider s'il faut
    reextraire un fichier lors d'une synchro.
    """
    st = path.stat()
    return hashlib.sha256(
        f"v{EXTRACTOR_VERSION}:{st.st_size}:{st.st_mtime_ns}".encode()
    ).hexdigest()


def file_to_document(path: Path) -> dict[str, object] | None:
    try:
        text = extract_text(path)
    except Exception:
        return None
    if not text.strip():
        return None

    st = path.stat()
    indexed = f"Fichier local : {path.name}\nChemin : {path}\n\n{text}"
    return {
        "source": "file",
        "external_id": str(path.resolve()),
        "title": path.name,
        "author": None,
        "ts": int(st.st_mtime),
        "url": path.resolve().as_uri(),
        "body": indexed,
        "meta": {
            "folder": str(path.parent),
            "extension": path.suffix.lower(),
            "size_bytes": st.st_size,
        },
        "content_hash": stat_hash(path),
    }


def sync(conn: sqlite3.Connection, cfg: Config, **_: object) -> list[int]:
    """Indexe les dossiers configures et retire les fichiers disparus."""
    if not cfg.files.roots:
        return []

    seen: set[str] = set()
    changed: list[int] = []

    for path in iter_files(cfg.files):
        key = str(path.resolve())
        seen.add(key)
        row = conn.execute(
            "SELECT id, content_hash FROM documents WHERE source = 'file' AND external_id = ?",
            (key,),
        ).fetchone()
        try:
            if row and row["content_hash"] == stat_hash(path):
                continue  # inchange depuis la derniere indexation
        except OSError:
            continue

        # Extraction AVANT toute ecriture, puis commit immediat : la base n'est
        # verrouillee que le temps d'une insertion, pas pendant la lecture d'un PDF.
        # Sans cela, une synchro de 100 PDF bloque l'interface (chat compris) plusieurs
        # minutes avec "database is locked".
        doc = file_to_document(path)
        if doc is None:
            continue
        doc_id, modified = db.upsert_document(conn, **doc)
        conn.commit()
        if modified:
            changed.append(doc_id)

    # Fichiers supprimes ou deplaces depuis la derniere synchro.
    for row in conn.execute("SELECT id, external_id FROM documents WHERE source = 'file'"):
        if row["external_id"] not in seen and not Path(row["external_id"]).exists():
            db.delete_document(conn, row["id"])

    conn.commit()
    return changed
